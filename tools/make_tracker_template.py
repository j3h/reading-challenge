"""Build tracker_template.xlsx: the data-entry workbook Jamie fills in (upload it to Google Sheets).

    python tools/make_tracker_template.py [output.xlsx]

Tabs: Start here, Entries (one row per paper form), Prizes (one row per prize handed out),
Roster (one row per child, with computed totals), Find a child (partial-name search), Levels.
Only functions that work in both Excel and Google Sheets are used (no FILTER/XLOOKUP).
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ENTRY_ROWS = 3000
PRIZE_ROWS = 3000
ROSTER_ROWS = 600
LEVEL_ROWS = 30  # Levels!2:30
FIND_RESULTS = 25
MONTHS = ",".join(f"{y}-{m:02d}" for y, m in
                  [(2026, 9), (2026, 10), (2026, 11), (2026, 12), (2027, 1),
                   (2027, 2), (2027, 3), (2027, 4), (2027, 5), (2027, 6)])

LEVELS = [
    (1, "First Pages", "Bookmark"),
    (100, "100 Pages", "Pencil and Stickers"),
    (500, "500 Pages", "Ribbon"),
    (1000, "1,000 Pages", "BCS Waterbottle"),
    (2500, "2,500 Pages", "Pen"),
    (5000, "5,000 Pages", "Metal Bookmark"),
    (7500, "7,500 Pages", "Baskin Robbin’s Ice Cream Scoop Giftcard"),
    (10000, "10,000 Pages", "BCS T-Shirt"),
    (12500, "12,500 Pages", "$5 Children’s Place Giftcard"),
    (15000, "15,000 Pages", "Fidget Toy"),
    (20000, "20,000 Pages", "Stickers"),
    (25000, "25,000 Pages", "Enamel Pin"),
    (30000, "30,000 Pages", "$5 Puddletown Games Giftcard"),
    (40000, "40,000 Pages", "Book"),
    (55000, "55,000 Pages", "Medal"),
    (70000, "70,000 Pages", "Engraved Bookmark"),
    (85000, "85,000 Pages", "$10 Powell’s Giftcard"),
    (90000, "90,000 Pages", "Lunch w/Principal Presely"),
    (100000, "100,000 Pages", "Trophy"),
]

FONT = "Arial"
NAVY = "1F3A5F"
f_body = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_note = Font(name=FONT, size=9, italic=True, color="666666")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_calc = PatternFill("solid", fgColor="EEF1F5")
fill_find = PatternFill("solid", fgColor="FFF6CC")
fill_warn = PatternFill("solid", fgColor="FFE3A3")
fill_good = PatternFill("solid", fgColor="CDEBC9")
thin = Side(style="thin", color="D5DAE0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, row, labels, widths):
    for i, (label, width) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.border = f_head, fill_head, border
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 32


def style_body(ws, first_row, last_row, ncols, calc_cols=()):
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.border = f_body, border
            if c in calc_cols:
                cell.fill = fill_calc


def dropdown(ws, formula, cell_range, allow_blank=True, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank,
                        showErrorMessage=strict, errorStyle="stop" if strict else "warning")
    ws.add_data_validation(dv)
    dv.add(cell_range)


wb = Workbook()

# ---------------------------------------------------------------- Start here
ws = wb.active
ws.title = "Start here"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
for col, w in zip("BCDEFGH", (22, 12, 10, 22, 18, 26, 14)):
    ws.column_dimensions[col].width = w
ws["B2"] = "Reading Challenge tracker"
ws["B2"].font = f_title
lines = [
    ("How this works", True),
    ("Type one row on Entries for every paper form, the day it arrives, even if you can't tell who the child is.", False),
    ("Gray columns are formulas: don't type in them. White columns are for you.", False),
    ("", False),
    ("Entering a form (Entries tab)", True),
    ("1. Fill in the month, pages, and the name and teacher exactly as written on the form.", False),
    ("2. If you know the child, pick them in the “Child” column (type a few letters to narrow the list). "
     "Leave it blank if you don't know yet.", False),
    ("3. Not sure who “Reid” is? Use the Find a child tab: type part of a name or teacher and it lists every match. "
     "Come back later and fill in the Child column.", False),
    ("4. Rows whose Status says “Needs child” are orange. Their pages are NOT counted until you pick a child.", False),
    ("5. A child's new kid isn't in the list? Add them on the Roster tab first, then pick them.", False),
    ("", False),
    ("Fixing mistakes", True),
    ("Don't overwrite a wrong number. Add a new row at the bottom for the same child and month: the lowest row wins.", False),
    ("A duplicate or junk form: put x in the Void column. Don't delete rows and don't sort the Entries tab "
     "(use Data > Create a filter view if you want to sort).", False),
    ("", False),
    ("Handing out prizes (Prizes tab)", True),
    ("Add a row when a prize is handed over: pick the child and the level. The Roster tab then shows who is still owed.", False),
    ("", False),
    ("Privacy", True),
    ("This workbook holds children's names. Keep it in the school's Google account, shared only with people who need it.", False),
    ("", False),
    ("Example (what a filled-in row looks like)", True),
]
r = 4
for text, bold in lines:
    c = ws.cell(row=r, column=2, value=text)
    c.font = f_bold if bold else f_body
    if text and not bold:
        c.alignment = Alignment(wrap_text=False)
    r += 1
ex_head = ["Entered at", "Month", "Pages", "Name on form", "Teacher on form", "Child", "Void"]
for i, label in enumerate(ex_head):
    c = ws.cell(row=r, column=2 + i, value=label)
    c.font, c.fill, c.border = f_head, fill_head, border
ex = ["2026-10-02 14:05", "2026-10", 140, "Reid", "Lopez", "(pick later)", ""]
for i, v in enumerate(ex):
    c = ws.cell(row=r + 1, column=2 + i, value=v)
    c.font, c.border = f_body, border
ws.cell(row=r + 3, column=2, value="Tip: in Google Sheets, Ctrl+Alt+Shift+; types the current date and time.").font = f_note

# ---------------------------------------------------------------- Levels
wl = wb.create_sheet("Levels")
header(wl, 1, ["Threshold (pages)", "Level", "Prize"], [18, 18, 46])
for i, (t, label, prize) in enumerate(LEVELS, start=2):
    wl.cell(row=i, column=1, value=t).number_format = "#,##0"
    wl.cell(row=i, column=2, value=label)
    wl.cell(row=i, column=3, value=prize)
style_body(wl, 2, LEVEL_ROWS, 3)
wl.cell(row=LEVEL_ROWS + 2, column=1,
        value="Copy of levels.csv. Keep the two in sync; the certificate program reads levels.csv.").font = f_note
wl.freeze_panes = "A2"

# ---------------------------------------------------------------- Roster
wr = wb.create_sheet("Roster")
header(wr, 1,
       ["Child ID", "Name", "Teacher", "School year", "Pick-list name", "Total pages",
        "Last prize level given", "Highest prize owed", "Levels owed", "Search helper"],
       [10, 26, 18, 12, 36, 12, 14, 18, 10, 10])
for i in range(2, ROSTER_ROWS + 1):
    wr.cell(row=i, column=1, value=f"C{i - 1:04d}")
    wr.cell(row=i, column=5, value=(
        f'=IF(B{i}="","",IF(COUNTIFS($B$2:$B${ROSTER_ROWS},B{i},$C$2:$C${ROSTER_ROWS},C{i})>1,'
        f'B{i}&" ("&C{i}&") "&A{i},B{i}&" ("&C{i}&")"))'))
    wr.cell(row=i, column=6, value=(
        f'=IF(B{i}="","",SUMIFS(Entries!$J$2:$J${ENTRY_ROWS},Entries!$H$2:$H${ENTRY_ROWS},A{i}))'))
    wr.cell(row=i, column=7, value=(
        f'=IF(B{i}="","",_xlfn.MAXIFS(Prizes!$D$2:$D${PRIZE_ROWS},Prizes!$E$2:$E${PRIZE_ROWS},A{i},'
        f'Prizes!$G$2:$G${PRIZE_ROWS},""))'))
    wr.cell(row=i, column=8, value=(
        f'=IF(B{i}="","",IF(_xlfn.MAXIFS(Levels!$A$2:$A${LEVEL_ROWS},Levels!$A$2:$A${LEVEL_ROWS},"<="&F{i})>G{i},'
        f'INDEX(Levels!$B$2:$B${LEVEL_ROWS},MATCH(_xlfn.MAXIFS(Levels!$A$2:$A${LEVEL_ROWS},'
        f'Levels!$A$2:$A${LEVEL_ROWS},"<="&F{i}),Levels!$A$2:$A${LEVEL_ROWS},0)),""))'))
    wr.cell(row=i, column=9, value=(
        f'=IF(B{i}="","",COUNTIFS(Levels!$A$2:$A${LEVEL_ROWS},"<="&F{i},Levels!$A$2:$A${LEVEL_ROWS},">"&G{i}))'))
    wr.cell(row=i, column=10, value=(
        f'=IF(OR(B{i}="",TRIM(\'Find a child\'!$B$3)=""),"",IF(ISNUMBER(SEARCH(TRIM(\'Find a child\'!$B$3),'
        f'B{i}&" "&C{i})),COUNT(J$1:J{i - 1})+1,""))'))
style_body(wr, 2, ROSTER_ROWS, 10, calc_cols=(1, 5, 6, 7, 8, 9, 10))
for i in range(2, ROSTER_ROWS + 1):
    wr.cell(row=i, column=6).number_format = "#,##0"
    wr.cell(row=i, column=7).number_format = "#,##0"
wr.conditional_formatting.add(
    f"H2:I{ROSTER_ROWS}", FormulaRule(formula=['AND($I2<>"",$I2>0)'], fill=fill_good))
wr.freeze_panes = "C2"

# ---------------------------------------------------------------- Entries
we = wb.create_sheet("Entries", 1)
header(we, 1,
       ["Entered at", "Month", "Pages", "Name on form", "Teacher on form", "Child", "Void (x)",
        "Child ID", "Status", "Pages counted"],
       [18, 10, 9, 24, 18, 36, 8, 10, 18, 10])
for i in range(2, ENTRY_ROWS + 1):
    we.cell(row=i, column=2).number_format = "@"
    we.cell(row=i, column=8, value=(
        f'=IF(F{i}="","",IFERROR(INDEX(Roster!$A$2:$A${ROSTER_ROWS},'
        f'MATCH(F{i},Roster!$E$2:$E${ROSTER_ROWS},0)),""))'))
    we.cell(row=i, column=9, value=(
        f'=IF(AND(D{i}="",C{i}="",F{i}=""),"",IF(G{i}<>"","Void",IF(F{i}="","Needs child",'
        f'IF(H{i}="","Child not found",IF(OR(B{i}="",NOT(ISNUMBER(C{i}))),"Needs month/pages","OK")))))'))
    we.cell(row=i, column=10, value=(
        f'=IF(I{i}="","",IF(I{i}="OK",IF(COUNTIFS(H{i + 1}:H${ENTRY_ROWS + 1},H{i},B{i + 1}:B${ENTRY_ROWS + 1},B{i},'
        f'G{i + 1}:G${ENTRY_ROWS + 1},"")=0,C{i},0),0))'))
style_body(we, 2, ENTRY_ROWS, 10, calc_cols=(8, 9, 10))
we.conditional_formatting.add(
    f"A2:J{ENTRY_ROWS}",
    FormulaRule(formula=['OR($I2="Needs child",$I2="Child not found",$I2="Needs month/pages")'], fill=fill_warn))
we.conditional_formatting.add(
    f"A2:J{ENTRY_ROWS}",
    FormulaRule(formula=['$I2="Void"'], font=Font(name=FONT, color="999999", strike=True)))
dropdown(we, f'"{MONTHS}"', f"B2:B{ENTRY_ROWS}")
dropdown(we, f"=Roster!$E$2:$E${ROSTER_ROWS}", f"F2:F{ENTRY_ROWS}")
dropdown(we, '"x"', f"G2:G{ENTRY_ROWS}")
we.freeze_panes = "A2"
we.cell(row=1, column=12, value=("To add a month beyond June 2027, edit the Month dropdown "
                                 "(Data > Data validation).")).font = f_note

# ---------------------------------------------------------------- Prizes
wp = wb.create_sheet("Prizes", 2)
header(wp, 1,
       ["Given on", "Child", "Level", "Threshold", "Child ID", "Prize", "Void (x)", "Status"],
       [14, 36, 16, 11, 10, 38, 8, 18])
for i in range(2, PRIZE_ROWS + 1):
    wp.cell(row=i, column=4, value=(
        f'=IF(C{i}="","",IFERROR(INDEX(Levels!$A$2:$A${LEVEL_ROWS},MATCH(C{i},Levels!$B$2:$B${LEVEL_ROWS},0)),""))'))
    wp.cell(row=i, column=5, value=(
        f'=IF(B{i}="","",IFERROR(INDEX(Roster!$A$2:$A${ROSTER_ROWS},'
        f'MATCH(B{i},Roster!$E$2:$E${ROSTER_ROWS},0)),""))'))
    wp.cell(row=i, column=6, value=(
        f'=IF(C{i}="","",IFERROR(INDEX(Levels!$C$2:$C${LEVEL_ROWS},MATCH(C{i},Levels!$B$2:$B${LEVEL_ROWS},0)),""))'))
    wp.cell(row=i, column=8, value=(
        f'=IF(AND(B{i}="",C{i}=""),"",IF(G{i}<>"","Void",IF(OR(B{i}="",E{i}=""),"Needs child",'
        f'IF(D{i}="","Needs level","OK"))))'))
    wp.cell(row=i, column=4).number_format = "#,##0"
style_body(wp, 2, PRIZE_ROWS, 8, calc_cols=(4, 5, 6, 8))
wp.conditional_formatting.add(
    f"A2:H{PRIZE_ROWS}",
    FormulaRule(formula=['OR($H2="Needs child",$H2="Needs level")'], fill=fill_warn))
wp.conditional_formatting.add(
    f"A2:H{PRIZE_ROWS}",
    FormulaRule(formula=['$H2="Void"'], font=Font(name=FONT, color="999999", strike=True)))
dropdown(wp, f"=Roster!$E$2:$E${ROSTER_ROWS}", f"B2:B{PRIZE_ROWS}")
dropdown(wp, f"=Levels!$B$2:$B${LEVEL_ROWS}", f"C2:C{PRIZE_ROWS}")
dropdown(wp, '"x"', f"G2:G{PRIZE_ROWS}")
wp.freeze_panes = "A2"

# ---------------------------------------------------------------- Find a child
wf = wb.create_sheet("Find a child", 3)
wf.sheet_view.showGridLines = False
wf["A1"] = "Find a child"
wf["A1"].font = f_title
wf["A3"] = "Type part of a name or teacher:"
wf["A3"].font = f_bold
wf["B3"].fill, wf["B3"].font, wf["B3"].border = fill_find, Font(name=FONT, size=12, bold=True), border
wf["A4"] = "Matches:"
wf["A4"].font = f_body
wf["B4"] = f"=COUNT(Roster!$J$2:$J${ROSTER_ROWS})"
wf["B4"].font = f_body
wf["B4"].alignment = Alignment(horizontal="left")
wf["C4"] = f'=IF(B4>{FIND_RESULTS},"Showing the first {FIND_RESULTS}; type more letters to narrow it down.","")'
wf["C4"].font = f_note
header(wf, 6, ["Name", "Teacher", "School year", "Total pages", "Last prize level given", "Highest prize owed"],
       [30, 22, 14, 14, 18, 20])
wf.column_dimensions["A"].width = 30
src = ["B", "C", "D", "F", "G", "H"]
for k in range(1, FIND_RESULTS + 1):
    row = 6 + k
    for j, col in enumerate(src, start=1):
        wf.cell(row=row, column=j, value=(
            f'=IFERROR(INDEX(Roster!${col}$2:${col}${ROSTER_ROWS},'
            f'MATCH({k},Roster!$J$2:$J${ROSTER_ROWS},0)),"")'))
style_body(wf, 7, 6 + FIND_RESULTS, 6, calc_cols=(1, 2, 3, 4, 5, 6))
for k in range(FIND_RESULTS):
    wf.cell(row=7 + k, column=4).number_format = "#,##0"
    wf.cell(row=7 + k, column=5).number_format = "#,##0"

# Tab order: Start here, Entries, Prizes, Find a child, Roster, Levels
wb.move_sheet("Roster", offset=len(wb.sheetnames))
wb.move_sheet("Levels", offset=len(wb.sheetnames))

out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "tracker_template.xlsx"
wb.save(out)
print(f"wrote {out}")
